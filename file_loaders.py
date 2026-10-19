"""
file_loaders.py - turn many file types into plain text for the RAG pipeline.

Each loader takes the raw bytes of a file and returns plain text. The rest of
the app (chunking, embeddings, FAISS, Gemini) doesn't care where the text came
from, so supporting a new file type only means adding a loader here.

This module has no Streamlit dependency, so it can be tested on its own.
"""

import io
import os
import sqlite3
import tempfile

import pandas as pd
from PyPDF2 import PdfReader

# Tabular files / databases can be huge. Embedding every row is slow and costs
# API quota, so only the first MAX_ROWS rows of each sheet/table are indexed.
MAX_ROWS = 2000

# Anything that is already plain text (docs, config, source code, scripts).
TEXT_EXTENSIONS = {
    # DOCUMENTS / MARKUP
    "txt", "md", "markdown", "rst", "log", "html", "htm", "xml",
    # DATA / CONGIF
    "json", "yaml", "yml", "toml", "ini", "cfg", "env",
    # PROGRAMMING LANGUAGES
    "py", "java", "php", "js", "ts", "jsx", "tsx", "c", "cpp", "h", "hpp",
    "cs", "go", "rb", "rs", "kt", "swift", "scala", "lua", "pl", "r",
    # SCRIPTS / QUERIES
    "sql", "sh", "bat", "ps1", "css",
}
SQLITE_EXTENSIONS = {"sqlite", "sqlite3", "db"}
SUPPORTED_EXTENSIONS = sorted(
    {"pdf", "csv", "xlsx", "docx", "pptx"} | TEXT_EXTENSIONS | SQLITE_EXTENSIONS
)


def _decode(data: bytes) -> str:
    """Decode bytes as text, trying common encodings before giving up."""
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")

def _dataframe_to_text(df: pd.DataFrame, label: str, notes: list) -> str:
    """
    Turn a table into 'column: value' lines, one per row. Embedding models
    understand "Name: Asha; City: Pune" far better than a bare "Asha,Pune".
    """
    total_rows = len(df)
    if total_rows > MAX_ROWS:
        notes.append(f"{label}: only the first {MAX_ROWS:,} of {total_rows:,} rows were indexed.")
        df = df.head(MAX_ROWS)

    lines = [f"Table: {label}", f"Columns: {', '.join(map(str, df.columns))}"]
    for _, row in df.iterrows():
        cells = [f"{col}: {val}" for col, val in row.items() if pd.notna(val)]
        if cells:
            lines.append("; ".join(cells))
    return "\n".join(lines)


# INDIVIDUAL LOADERS
def _load_pdf(data, name, notes):
    reader = PdfReader(io.BytesIO(data))
    pages = []
    for page in reader.pages:
        extracted = page.extract_text()
        if extracted:
            pages.append(extracted)
    return "\n".join(pages)

def _load_text(data, name, notes):
    return _decode(data)

def _load_csv(data, name, notes):
    df = pd.read_csv(io.StringIO(_decode(data)), sep=None, engine="python")
    return _dataframe_to_text(df, name, notes)

def _load_xlsx(data, name, notes):
    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)  # needs openpyxl
    parts = [_dataframe_to_text(df, f"{name} / sheet '{sheet}'", notes)
             for sheet, df in sheets.items()]
    return "\n\n".join(parts)

def _load_docx(data, name, notes):
    from docx import Document  # python-docx
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)

def _load_pptx(data, name, notes):
    from pptx import Presentation  # python-pptx
    prs = Presentation(io.BytesIO(data))
    parts = []
    for number, slide in enumerate(prs.slides, start=1):
        parts.append(f"--- Slide {number} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            parts.append("Speaker notes: " + slide.notes_slide.notes_text_frame.text)
    return "\n".join(parts)

def _load_sqlite(data, name, notes):
    # sqlite3 needs a real file on disk, so write the upload to a temp file.
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".sqlite")
    try:
        tmp.write(data)
        tmp.close()
        conn = sqlite3.connect(tmp.name)
        try:
            tables = conn.execute(
                "SELECT name, sql FROM sqlite_master "
                "WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            ).fetchall()
            parts = []
            for table_name, create_sql in tables:
                parts.append(f"Database table: {table_name}\nSchema: {create_sql}")
                df = pd.read_sql_query(f'SELECT * FROM "{table_name}"', conn)
                parts.append(_dataframe_to_text(df, f"{name} / table '{table_name}'", notes))
            return "\n\n".join(parts)
        finally:
            conn.close()
    finally:
        os.remove(tmp.name)


# PUBLIC ENTRY POINT
def load_file(name: str, data: bytes):
    """
    Returns (text, notes). `notes` is a list of human-readable warnings
    (e.g. truncated tables) the UI can show to the user.
    Raises ValueError for unsupported types.
    """
    ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
    notes = []
    if ext == "pdf":
        text = _load_pdf(data, name, notes)
    elif ext == "csv":
        text = _load_csv(data, name, notes)
    elif ext == "xlsx":
        text = _load_xlsx(data, name, notes)
    elif ext == "docx":
        text = _load_docx(data, name, notes)
    elif ext == "pptx":
        text = _load_pptx(data, name, notes)
    elif ext in SQLITE_EXTENSIONS:
        text = _load_sqlite(data, name, notes)
    elif ext in TEXT_EXTENSIONS:
        text = _load_text(data, name, notes)
    else:
        raise ValueError(f"Unsupported file type: .{ext}")

    return text, notes